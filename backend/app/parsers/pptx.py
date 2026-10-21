"""PowerPoint (.pptx) extraction: slide titles as headings, body text, tables, speaker notes."""

import io
from collections.abc import Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape

from app.parsers.base import ParseError


def _walk(shapes) -> Iterator[BaseShape]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def _text_lines(shape: BaseShape) -> list[str]:
    lines: list[str] = []
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        return lines
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip() or para.text.strip()
            if text:
                lines.append(f"{'  ' * para.level}- {text}" if para.level else text)
    return lines


def parse_pptx(data: bytes) -> str:
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("Could not open this .pptx file. It may be corrupt.") from exc

    out: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        block = [f"# {title or f'Slide {index}'}"]
        for shape in _walk(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            block.extend(_text_lines(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            note_text = notes.text.strip() if notes is not None else ""
            if note_text:
                block.append(f"Speaker notes: {note_text}")
        if len(block) > 1 or title:
            out.append("\n".join(block))
    return "\n\n".join(out)
