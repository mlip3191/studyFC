"""Regenerate the tiny fixture files: `uv run python tests/make_fixtures.py`."""

from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation

OUT = Path(__file__).parent / "fixtures"

MARKDOWN = """# Sample Guide

Intro paragraph.

## Photosynthesis

Plants convert light into chemical energy.

## Respiration

Cells release energy from glucose.
"""


def make_md() -> None:
    (OUT / "sample.md").write_text(MARKDOWN, encoding="utf-8")


def make_pdf() -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for text, size in [
        ("Sample Guide", 24),
        ("Photosynthesis", 16),
        ("Plants convert light into chemical energy.", 11),
        ("Chlorophyll absorbs mostly red and blue light.", 11),
        ("Respiration", 16),
        ("Cells release energy from glucose.", 11),
    ]:
        page.insert_text((72, y), text, fontsize=size)
        y += size + 14
    doc.save(OUT / "sample.pdf")
    doc.close()

    blank = pymupdf.open()
    blank.new_page()
    blank.save(OUT / "blank.pdf")
    blank.close()


def make_docx() -> None:
    doc = Document()
    doc.add_heading("Sample Guide", level=1)
    doc.add_paragraph("Intro paragraph.")
    doc.add_heading("Photosynthesis", level=2)
    doc.add_paragraph("Plants convert light into chemical energy.")
    doc.add_paragraph("Chlorophyll", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "ATP"
    table.cell(1, 1).text = "Energy currency"
    doc.save(OUT / "sample.docx")


def make_pptx() -> None:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Photosynthesis"
    s1.placeholders[1].text = "Light reactions\nCalvin cycle"
    s1.notes_slide.notes_text_frame.text = "Mention chloroplasts."
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Respiration"
    s2.placeholders[1].text = "Glycolysis"
    prs.save(OUT / "sample.pptx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_md()
    make_pdf()
    make_docx()
    make_pptx()
